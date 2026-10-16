"""Text aus Dokumenten – PDF, DOCX, Text, und Scans über OCR.

Ein PDF ohne Textschicht erkennt man daran, dass pypdf fast nichts findet.
Dann übernimmt der OCR-Dienst (Textract in AWS). Ist keiner eingerichtet,
bleibt das Dokument außen vor, mit Hinweis – lieber eine Lücke, die man
sieht, als eine Analyse, die ein Dokument still übergeht.

**Seite für Seite, nicht Dokument für Dokument – der Fehler vom 16.10.2026.**
Aus der Kanzlei: „Das Programm hat beim Einlesen ganze Blöcke nicht
mitgenommen und zwar an Stellen, die sehr wichtig waren."

Bis dahin entschied der **Mittelwert über alle Seiten**, ob ein PDF ein Scan
ist. Das geht genau bei dem Dokument schief, das in einer Akte am häufigsten
vorkommt: der digital geschriebene Schriftsatz mit der **eingescannten Anlage
dahinter**. Zweiundzwanzig Seiten Text und acht Seiten Scan ergeben einen
hohen Mittelwert; das Dokument galt als lesbar, pypdf gab für die acht Seiten
nichts zurück, und niemand erfuhr davon. Gezählt wurde nur, ob das **ganze**
Dokument null Zeichen hat.

Dass es die wichtigen Stellen traf, ist kein Zufall: Gescannt wird, was von
außen kommt und unterschrieben ist – das Attest, das Urteil, die
Zustellungsurkunde, der handschriftliche Vermerk.

Entschieden wird deshalb **je Seite**. Hat ein PDF einzelne leere Seiten,
gehen genau diese durch die Texterkennung (das kostet je Seite, also nur die,
die es braucht). Bleibt danach eine Seite ohne Text – weil keine
Texterkennung eingerichtet ist oder weil auch sie nichts fand –, trägt die
Seite einen **Hinweis** statt stillschweigend zu fehlen.
"""

from __future__ import annotations

import io
from pathlib import PurePosixPath

from jusperfect.domain.models import Dokument, Seite
from jusperfect.ocr.base import OcrDienst, OcrFehler, OcrNichtVerfuegbar

_SCAN_SCHWELLE_ZEICHEN_JE_SEITE = 30

#: Was an der Seite steht, die auch nach der Texterkennung leer blieb.
OHNE_TEXT = "Auf dieser Seite wurde kein Text gefunden (vermutlich ein Scan)."
OHNE_TEXTERKENNUNG = ("Diese Seite ist vermutlich ein Scan. Es ist keine "
                      "Texterkennung eingerichtet (JUSPERFECT_OCR=textract).")


class NichtUnterstuetzt(ValueError):
    pass


def _leere_seiten(seiten: list[Seite]) -> list[int]:
    """Die Seitennummern, auf denen so gut wie nichts steht.

    Dieselbe Schwelle wie früher, nur eben **je Seite** gemessen und nicht im
    Mittel über das Dokument. Ein paar Zeichen hat auch eine gescannte Seite
    manchmal - eine Kopfzeile aus der Fax-Software, eine Seitenzahl.
    """
    return [s.nummer for s in seiten
            if len(s.text.strip()) < _SCAN_SCHWELLE_ZEICHEN_JE_SEITE]


def _ist_scan(seiten: list[Seite]) -> bool:
    """Ist das GANZE Dokument ein Scan? Dann geht es als Ganzes durch die OCR."""
    if not seiten:
        return True
    return len(_leere_seiten(seiten)) == len(seiten)


def seiten_nennen(nummern: list[int]) -> str:
    """„3", „3 und 7", „3, 7 bis 9 und 14" - zusammenhängendes zusammengefasst.

    Eine Aufzählung von vierzehn Einzelzahlen liest niemand; „7 bis 20" sagt
    auf einen Blick, dass der ganze hintere Teil fehlt.
    """
    if not nummern:
        return ""
    gruppen: list[list[int]] = [[nummern[0]]]
    for nr in nummern[1:]:
        if nr == gruppen[-1][-1] + 1:
            gruppen[-1].append(nr)
        else:
            gruppen.append([nr])
    teile = [str(g[0]) if len(g) == 1 else f"{g[0]} bis {g[-1]}" for g in gruppen]
    if len(teile) == 1:
        return teile[0]
    return ", ".join(teile[:-1]) + " und " + teile[-1]


def _teil_pdf(daten: bytes, nummern: list[int]) -> bytes | None:
    """Nur diese Seiten als eigenes PDF - oder None, wenn das nicht geht.

    Die Texterkennung rechnet je Seite ab. Ein Schriftsatz mit drei
    gescannten Anlagenseiten soll drei Seiten kosten und nicht dreißig.
    """
    try:
        from pypdf import PdfReader, PdfWriter

        leser = PdfReader(io.BytesIO(daten))
        schreiber = PdfWriter()
        for nr in nummern:
            schreiber.add_page(leser.pages[nr - 1])
        puffer = io.BytesIO()
        schreiber.write(puffer)
        return puffer.getvalue()
    except Exception:
        return None


def _leere_seiten_erkennen(seiten: list[Seite], leer: list[int], daten: bytes,
                           ocr: OcrDienst, sitzung_id: str, dokument_id: str) -> tuple[list[Seite], bool]:
    """Die leeren Seiten eines sonst lesbaren PDFs nachträglich erkennen.

    Hier wird **nie** eine Ausnahme weitergereicht. Zweiundzwanzig gelesene
    Seiten wegzuwerfen, weil die acht gescannten nicht gehen, wäre genau der
    Schaden, der behoben werden soll. Was nicht geht, steht danach als Hinweis
    an der betroffenen Seite.
    """
    nach_nummer: dict[int, str] = {}
    grund = OHNE_TEXT
    benutzt = False
    teil = _teil_pdf(daten, leer)
    if teil is None:
        grund = ("Diese Seite ließ sich für die Texterkennung nicht aus dem "
                 "PDF herauslösen.")
    else:
        try:
            erkannt = ocr.erkennen(teil, "application/pdf", sitzung_id, dokument_id)
            benutzt = True
            # Die Texterkennung zählt im Teil-PDF wieder ab 1. Seite i darin
            # ist die Seite leer[i-1] im Original - sonst zeigte jede
            # Fundstelle auf die falsche Seite, und die Belegpflicht trüge nicht.
            texte = {s.nummer: s.text for s in erkannt}
            for stelle, nummer in enumerate(leer, start=1):
                nach_nummer[nummer] = texte.get(stelle, "")
        except OcrNichtVerfuegbar:
            grund = OHNE_TEXTERKENNUNG
        except OcrFehler as fehler:
            grund = f"Die Texterkennung dieser Seite schlug fehl: {fehler}"

    ergebnis: list[Seite] = []
    for s in seiten:
        if s.nummer not in leer:
            ergebnis.append(s)
            continue
        text = nach_nummer.get(s.nummer, "").strip()
        if text:
            ergebnis.append(Seite(nummer=s.nummer, text=text))
        else:
            ergebnis.append(Seite(nummer=s.nummer, text=s.text, hinweis=grund))
    return ergebnis, benutzt


def _pdf(daten: bytes) -> list[Seite]:
    from pypdf import PdfReader

    leser = PdfReader(io.BytesIO(daten))
    return [Seite(nummer=nr, text=(seite.extract_text() or "")) for nr, seite in enumerate(leser.pages, start=1)]


def _docx(daten: bytes) -> list[Seite]:
    import docx

    dokument = docx.Document(io.BytesIO(daten))
    teile = [a.text for a in dokument.paragraphs if a.text.strip()]
    for tabelle in dokument.tables:
        for zeile in tabelle.rows:
            zellen = [z.text.strip() for z in zeile.cells if z.text.strip()]
            if zellen:
                teile.append(" | ".join(zellen))
    return [Seite(nummer=1, text="\n\n".join(teile))]


def _text(daten: bytes) -> list[Seite]:
    try:
        text = daten.decode("utf-8")
    except UnicodeDecodeError:
        text = daten.decode("latin-1")
    return [Seite(nummer=1, text=text)]


# ── Weitere Formate ──────────────────────────────────────────────────────────
#
# Aus der Kanzlei am 25.09.2026: "24 von 38 Dokumenten gelesen, 14 nicht: 14 in
# einem Format, das nicht gelesen werden kann." Und dann: "ja bitte, er soll
# alle Dokumenttypen lesen koennen."
#
# In einer Akte liegen nicht nur Schriftsaetze. Es liegen E-Mails darin, alte
# Word-Dateien, Tabellen, ausgedruckte Webseiten. Ein Programm, das davon die
# Haelfte nicht ansieht, hat die Akte nicht gelesen - und niemand merkt es,
# weil eine kurze Antwort wie eine knappe Akte aussieht.
#
# Gebaut wird mit dem, was Python mitbringt, wo es geht: Eine Abhaengigkeit
# weniger ist eine Fehlerquelle weniger auf einem Rechner ohne Internet. Nur
# fuer .msg (Outlook) gibt es keinen Weg ohne Paket - das Format ist ein
# eigener Dateibaum, kein Text.


def _eml(daten: bytes) -> list[Seite]:
    """Eine E-Mail: Kopfzeilen und Text, Anhaenge nur dem Namen nach.

    Wer eine Mail in die Akte legt, meint meist ihren Inhalt - aber die
    Kopfzeilen tragen die halbe Aussage: Von wem, an wen, wann. Ohne sie
    liesse sich ein Satz aus der Mail nicht zuordnen.
    """
    from email import message_from_bytes
    from email.header import decode_header, make_header
    from email.policy import default as standard

    nachricht = message_from_bytes(daten, policy=standard)

    def lesbar(wert: object) -> str:
        if not wert:
            return ""
        try:
            return str(make_header(decode_header(str(wert))))
        except Exception:
            return str(wert)

    kopf = []
    for feld, beschriftung in (("From", "Von"), ("To", "An"), ("Cc", "Kopie"),
                               ("Date", "Datum"), ("Subject", "Betreff")):
        wert = lesbar(nachricht.get(feld))
        if wert:
            kopf.append(f"{beschriftung}: {wert}")

    text = ""
    anhaenge: list[str] = []
    if nachricht.is_multipart():
        for teil in nachricht.walk():
            art = teil.get_content_type()
            name = teil.get_filename()
            if name:
                anhaenge.append(lesbar(name))
                continue
            if art == "text/plain" and not text:
                text = teil.get_content()
            elif art == "text/html" and not text:
                text = _html_zu_text(teil.get_content())
    else:
        inhalt = nachricht.get_content()
        text = (_html_zu_text(inhalt)
                if nachricht.get_content_type() == "text/html" else inhalt)

    teile = ["\n".join(kopf)] if kopf else []
    if text and text.strip():
        teile.append(text.strip())
    if anhaenge:
        # Die Anhaenge liegen in jur|nodes als eigene Dateien; hier steht nur,
        # DASS es sie gab - sonst fehlt im Zusammenhang ein Stueck.
        teile.append("Anhänge: " + ", ".join(anhaenge))
    return [Seite(nummer=1, text="\n\n".join(teile))]


def _msg(daten: bytes) -> list[Seite]:
    """Eine Outlook-Nachricht. Braucht als einzige ein eigenes Paket."""
    import tempfile

    import extract_msg

    # extract_msg liest von einem Pfad. Die Datei bleibt nur, solange sie
    # gelesen wird - Aktentext gehoert nicht laenger als noetig auf die Platte.
    with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as ablage:
        ablage.write(daten)
        pfad = ablage.name
    try:
        nachricht = extract_msg.Message(pfad)
        kopf = []
        for wert, beschriftung in ((nachricht.sender, "Von"),
                                   (nachricht.to, "An"),
                                   (nachricht.cc, "Kopie"),
                                   (nachricht.date, "Datum"),
                                   (nachricht.subject, "Betreff")):
            if wert:
                kopf.append(f"{beschriftung}: {wert}")
        teile = ["\n".join(kopf)] if kopf else []
        koerper = nachricht.body or ""
        if not koerper.strip() and getattr(nachricht, "htmlBody", None):
            roh = nachricht.htmlBody
            koerper = _html_zu_text(roh.decode("utf-8", "replace")
                                    if isinstance(roh, bytes) else roh)
        if koerper.strip():
            teile.append(koerper.strip())
        namen = [a.longFilename or a.shortFilename
                 for a in (nachricht.attachments or [])]
        namen = [n for n in namen if n]
        if namen:
            teile.append("Anhänge: " + ", ".join(namen))
        nachricht.close()
        return [Seite(nummer=1, text="\n\n".join(teile))]
    finally:
        import os
        try:
            os.unlink(pfad)
        except OSError:
            pass


def _html_zu_text(roh: str) -> str:
    """Aus einer Webseite oder HTML-Mail den Text - ohne Auszeichnung."""
    from html.parser import HTMLParser

    class Leser(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.teile: list[str] = []
            self.still = 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style", "head"):
                self.still += 1
            elif tag in ("p", "br", "div", "tr", "li", "h1", "h2", "h3", "h4"):
                self.teile.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style", "head") and self.still:
                self.still -= 1

        def handle_data(self, daten):
            if not self.still and daten.strip():
                self.teile.append(daten)

    leser = Leser()
    leser.feed(roh)
    zeilen = [z.strip() for z in "".join(leser.teile).splitlines()]
    return "\n".join(z for z in zeilen if z)


def _html(daten: bytes) -> list[Seite]:
    roh = daten.decode("utf-8", "replace")
    return [Seite(nummer=1, text=_html_zu_text(roh))]


def _rtf(daten: bytes) -> list[Seite]:
    """Rich Text - Text mit geschweiften Klammern drumherum.

    Kein vollstaendiger Leser, sondern einer, der die Steuerworte wegnimmt und
    den Text stehenlaesst. Fuer einen Schriftsatz genuegt das; Tabellen
    verlieren ihre Form, nicht ihren Inhalt.
    """
    import re

    roh = daten.decode("latin-1", "replace")
    # Ganze Bloecke, die nie Text sind (Schriftarten, Farben, Stile).
    roh = re.sub(r"\\\*\\[a-z]+(?:-?\d+)?[^{}]*", " ", roh)
    for block in ("fonttbl", "colortbl", "stylesheet", "info", "pict"):
        roh = re.sub(r"\{\\" + block + r"[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", " ", roh)
    # Zeichen, die als Nummer geschrieben sind.
    roh = re.sub(r"\\'([0-9a-fA-F]{2})",
                 lambda t: bytes([int(t.group(1), 16)]).decode("latin-1"), roh)
    roh = re.sub(r"\\u(-?\d+)\s?\??",
                 lambda t: chr(int(t.group(1)) % 65536), roh)
    # Absatzmarken vor dem Rest der Steuerworte.
    roh = re.sub(r"\\(?:par|line|page)\b", "\n", roh)
    roh = re.sub(r"\\tab\b", "\t", roh)
    roh = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", roh)
    roh = roh.replace("{", " ").replace("}", " ").replace("\\", " ")
    zeilen = [" ".join(z.split()) for z in roh.splitlines()]
    return [Seite(nummer=1, text="\n".join(z for z in zeilen if z))]


def _xml_text(roh: bytes, ende_zeichen: tuple[str, ...] = ()) -> str:
    """Aller Text aus einer XML-Datei, in der Reihenfolge des Dokuments."""
    import xml.etree.ElementTree as ET

    try:
        wurzel = ET.fromstring(roh)
    except ET.ParseError:
        return ""
    teile: list[str] = []
    for element in wurzel.iter():
        marke = element.tag.rsplit("}", 1)[-1]
        if marke in ende_zeichen:
            teile.append("\n")
            continue
        if element.text and element.text.strip():
            teile.append(element.text)
    return "".join(teile)


def _xlsx(daten: bytes) -> list[Seite]:
    """Eine Tabelle - jedes Blatt eine Seite.

    Ein xlsx ist ein ZIP mit XML darin; dafuer braucht es kein Paket. Gelesen
    werden die Zeichenketten und die Zellen in ihrer Reihenfolge. Formeln
    bleiben aussen vor - in einer Akte zaehlt, was dasteht.
    """
    import re
    import xml.etree.ElementTree as ET
    import zipfile

    with zipfile.ZipFile(io.BytesIO(daten)) as behaelter:
        namen = behaelter.namelist()
        gemeinsam: list[str] = []
        if "xl/sharedStrings.xml" in namen:
            try:
                wurzel = ET.fromstring(behaelter.read("xl/sharedStrings.xml"))
                for eintrag in wurzel:
                    gemeinsam.append("".join(
                        t.text or "" for t in eintrag.iter()
                        if t.tag.rsplit("}", 1)[-1] == "t"))
            except ET.ParseError:
                pass

        seiten: list[Seite] = []
        blaetter = sorted(n for n in namen
                          if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))
        for nummer, blatt in enumerate(blaetter, start=1):
            try:
                wurzel = ET.fromstring(behaelter.read(blatt))
            except ET.ParseError:
                continue
            zeilen: list[str] = []
            for zeile in wurzel.iter():
                if zeile.tag.rsplit("}", 1)[-1] != "row":
                    continue
                zellen: list[str] = []
                for zelle in zeile:
                    art = zelle.get("t")
                    wert = ""
                    for kind in zelle:
                        marke = kind.tag.rsplit("}", 1)[-1]
                        if marke == "v":
                            wert = kind.text or ""
                        elif marke == "is":
                            wert = "".join(t.text or "" for t in kind.iter()
                                           if t.tag.rsplit("}", 1)[-1] == "t")
                    if art == "s" and wert.isdigit() and int(wert) < len(gemeinsam):
                        wert = gemeinsam[int(wert)]
                    if wert.strip():
                        zellen.append(wert.strip())
                if zellen:
                    zeilen.append(" | ".join(zellen))
            if zeilen:
                seiten.append(Seite(nummer=nummer, text="\n".join(zeilen)))
        return seiten or [Seite(nummer=1, text="")]


def _pptx(daten: bytes) -> list[Seite]:
    """Eine Folienmappe - jede Folie eine Seite."""
    import re
    import zipfile

    with zipfile.ZipFile(io.BytesIO(daten)) as behaelter:
        folien = sorted(
            (n for n in behaelter.namelist()
             if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)))
        seiten = []
        for nummer, folie in enumerate(folien, start=1):
            text = _xml_text(behaelter.read(folie), ende_zeichen=("p",))
            zeilen = [" ".join(z.split()) for z in text.splitlines()]
            seiten.append(Seite(nummer=nummer,
                                text="\n".join(z for z in zeilen if z)))
        return seiten or [Seite(nummer=1, text="")]


def _art(dokument: Dokument) -> str:
    """Welcher Leser fuer dieses Dokument.

    Die Endung entscheidet, wo sie eindeutig ist. Der MIME-Typ steht daneben,
    weil jur|nodes ihn nicht immer mitgibt - dann steht dort
    "application/octet-stream", und ohne Endung waere jedes Dokument
    unbekannt.
    """
    mime = dokument.mime_typ.lower()
    endung = PurePosixPath(dokument.dateiname).suffix.lower()
    if mime == "application/pdf" or endung == ".pdf":
        return "pdf"
    if "wordprocessingml" in mime or endung in (".docx", ".docm", ".dotx"):
        return "docx"
    # E-Mails. In einer Akte liegen oft mehr davon als Schriftsaetze.
    if endung == ".msg" or "outlook" in mime:
        return "msg"
    if endung == ".eml" or mime in ("message/rfc822", "application/mbox"):
        return "eml"
    if endung in (".htm", ".html", ".xhtml") or mime in ("text/html", "application/xhtml+xml"):
        return "html"
    if endung == ".rtf" or "rtf" in mime:
        return "rtf"
    if endung in (".xlsx", ".xlsm") or "spreadsheetml" in mime:
        return "xlsx"
    if endung in (".pptx", ".pptm") or "presentationml" in mime:
        return "pptx"
    if mime.startswith("text/") or endung in (".txt", ".md", ".csv", ".log", ".json", ".xml"):
        return "text"
    if mime.startswith("image/") or endung in (".png", ".jpg", ".jpeg", ".tif",
                                               ".tiff", ".bmp", ".gif"):
        return "bild"
    # Die alten Microsoft-Formate. Sie lassen sich nicht zuverlaessig lesen,
    # und ein halb gelesener Schriftsatz ist schlimmer als ein fehlender:
    # Was daraus zitiert wuerde, stuende so nicht in der Akte.
    if endung in (".doc", ".xls", ".ppt"):
        return "alt"
    return "unbekannt"


def text_extrahieren(dokument: Dokument, daten: bytes, ocr: OcrDienst, sitzung_id: str) -> tuple[list[Seite], bool]:
    """Gibt die Seiten zurück und ob OCR verwendet wurde."""
    art = _art(dokument)
    if art == "pdf":
        try:
            seiten = _pdf(daten)
        except Exception as e:
            raise NichtUnterstuetzt(f"PDF nicht lesbar ({type(e).__name__}).") from e
        leer = _leere_seiten(seiten)
        if not leer:
            return seiten, False
        if len(leer) == len(seiten):
            # Ein durchgehender Scan: als Ganzes durch die Texterkennung. Geht
            # das nicht, bleibt das Dokument außen vor - mit Hinweis, wie bisher.
            return ocr.erkennen(daten, "application/pdf", sitzung_id, dokument.id), True
        # Gemischt - der häufige Fall: Schriftsatz digital, Anlage gescannt.
        return _leere_seiten_erkennen(seiten, leer, daten, ocr, sitzung_id, dokument.id)
    if art == "docx":
        try:
            return _docx(daten), False
        except Exception as e:
            raise NichtUnterstuetzt(f"Word-Dokument nicht lesbar ({type(e).__name__}).") from e
    if art == "text":
        return _text(daten), False
    if art == "bild":
        mime = dokument.mime_typ if dokument.mime_typ.startswith("image/") else "image/png"
        return ocr.erkennen(daten, mime, sitzung_id, dokument.id), True
    # Die uebrigen Formate. Jedes einzeln, damit die Meldung sagt, WAS nicht
    # ging - "nicht lesbar" allein hilft niemandem weiter.
    for kennung, leser, name in (("msg", _msg, "Outlook-Nachricht"),
                                 ("eml", _eml, "E-Mail"),
                                 ("html", _html, "Webseite"),
                                 ("rtf", _rtf, "RTF-Dokument"),
                                 ("xlsx", _xlsx, "Tabelle"),
                                 ("pptx", _pptx, "Folienmappe")):
        if art != kennung:
            continue
        try:
            return leser(daten), False
        except ImportError as e:
            raise NichtUnterstuetzt(
                f"Für Outlook-Nachrichten fehlt ein Baustein ({e.name}). "
                "Er wird beim nächsten Start mitinstalliert.") from e
        except Exception as e:
            raise NichtUnterstuetzt(f"{name} nicht lesbar ({type(e).__name__}).") from e
    if art == "alt":
        raise NichtUnterstuetzt(
            f"Altes Microsoft-Format ({PurePosixPath(dokument.dateiname).suffix}). "
            "Es lässt sich nicht zuverlässig lesen; in jurnodes als PDF oder "
            "als neueres Format ablegen.")
    raise NichtUnterstuetzt(f"Dateityp nicht unterstützt: {dokument.dateiname}")
