import io
import unittest

from jusperfect.domain.models import Dokument, Seite
from jusperfect.ocr.base import KeinOcr, OcrNichtVerfuegbar
from jusperfect.services.extraktion import NichtUnterstuetzt, text_extrahieren


class ZaehlendesOcr:
    def __init__(self):
        self.aufrufe = 0

    def erkennen(self, daten, mime, sitzung_id, dokument_id):
        self.aufrufe += 1
        return [Seite(nummer=1, text="OCR-Text")]


def _dok(name, mime="application/octet-stream"):
    return Dokument(id="D", akte_id="A", titel=name, dateiname=name, mime_typ=mime)


def _pdf_mit_text():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    w = PdfWriter()
    seite = w.add_blank_page(width=200, height=200)
    schrift = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                NameObject("/BaseFont"): NameObject("/Helvetica")})
    seite[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): w._add_object(schrift)})})
    inhalt = DecodedStreamObject()
    inhalt.set_data(b"BT /F1 12 Tf 20 100 Td (Unfall am 13.03.2026 auf der B 27 mit erheblichem Sachschaden am Fahrzeug) Tj ET")
    seite[NameObject("/Contents")] = w._add_object(inhalt)
    puffer = io.BytesIO()
    w.write(puffer)
    return puffer.getvalue()


def _pdf_leer():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    puffer = io.BytesIO()
    w.write(puffer)
    return puffer.getvalue()


class Extraktion(unittest.TestCase):
    def test_text(self):
        seiten, ocr = text_extrahieren(_dok("a.txt", "text/plain"), "Hallo Welt".encode("utf-8"), KeinOcr(), "s")
        self.assertEqual((seiten[0].text, ocr), ("Hallo Welt", False))

    def test_latin1_rueckfall(self):
        seiten, _ = text_extrahieren(_dok("a.txt"), "Testmann".encode("latin-1"), KeinOcr(), "s")
        self.assertEqual(seiten[0].text, "Testmann")

    def test_pdf_mit_textschicht(self):
        ocr = ZaehlendesOcr()
        seiten, verwendet = text_extrahieren(_dok("a.pdf", "application/pdf"), _pdf_mit_text(), ocr, "s")
        self.assertIn("Unfall am 13.03.2026", seiten[0].text)
        self.assertEqual((verwendet, ocr.aufrufe), (False, 0))

    def test_scan_geht_an_ocr(self):
        ocr = ZaehlendesOcr()
        seiten, verwendet = text_extrahieren(_dok("scan.pdf", "application/pdf"), _pdf_leer(), ocr, "s")
        self.assertEqual((seiten[0].text, verwendet, ocr.aufrufe), ("OCR-Text", True, 1))

    def test_scan_ohne_ocr(self):
        with self.assertRaises(OcrNichtVerfuegbar):
            text_extrahieren(_dok("scan.pdf", "application/pdf"), _pdf_leer(), KeinOcr(), "s")

    def test_docx(self):
        import docx

        d = docx.Document()
        d.add_paragraph("Absatz eins")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Betrag", "850,00 EUR"
        puffer = io.BytesIO()
        d.save(puffer)
        seiten, _ = text_extrahieren(_dok("b.docx"), puffer.getvalue(), KeinOcr(), "s")
        self.assertIn("Absatz eins", seiten[0].text)
        self.assertIn("Betrag | 850,00 EUR", seiten[0].text)

    def test_unbekannt(self):
        with self.assertRaises(NichtUnterstuetzt):
            text_extrahieren(_dok("x.xyz"), b"...", KeinOcr(), "s")
